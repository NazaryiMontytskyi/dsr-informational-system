"""
Заповнення зразків наказів (app/order_templates/*.docx) реальними даними.

Підхід:
1. "Прості" плейсхолдери (`<...>`) підставляються по тексту абзацу/комірки
   таблиці -- значення може бути одним рядком (та сама підстановка всюди,
   де зустрічається плейсхолдер) або списком рядків (кожне наступне
   входження того самого плейсхолдера в документі споживає наступне
   значення зі списку -- потрібно, наприклад, для "<Ім'я ПРІЗВИЩЕ>", яке в
   зразку наказу повторюється п'ять разів із різними людьми).
2. Пункт "НАКАЗУЮ:" -- це нумерований список Word, кількість підпунктів
   якого залежить від даних (наприклад, кількість керівників гуртка), тому
   він переформовується повністю: перший підпункт клонується стільки
   разів, скільки потрібно пунктів, і в кожен підставляється текст.
"""

from __future__ import annotations

import copy
import re
from pathlib import Path

import docx
from docx.document import Document as DocumentObject
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

PLACEHOLDER_FONT = "Times New Roman"


def _iter_block_items(parent):
    if isinstance(parent, DocumentObject):
        parent_elm = parent.element.body
    else:
        parent_elm = parent._tc
    for child in parent_elm.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            yield Table(child, parent)


def _iter_all_paragraphs(doc: DocumentObject):
    for block in _iter_block_items(doc):
        if isinstance(block, Paragraph):
            yield block
        else:
            for row in block.rows:
                for cell in row.cells:
                    yield from _iter_all_paragraphs_in_cell(cell)


def _iter_all_paragraphs_in_cell(cell):
    for block in _iter_block_items(cell):
        if isinstance(block, Paragraph):
            yield block
        else:
            for row in block.rows:
                for c in row.cells:
                    yield from _iter_all_paragraphs_in_cell(c)


def _set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    runs = paragraph.runs
    if not runs:
        run = paragraph.add_run(text)
        run.font.name = PLACEHOLDER_FONT
        return
    runs[0].text = text
    runs[0].font.name = PLACEHOLDER_FONT
    for r in runs[1:]:
        r.text = ""


def replace_simple_placeholders(doc: DocumentObject, replacements: dict[str, str | list[str]]) -> None:
    counters = {k: 0 for k in replacements}
    pattern = re.compile("|".join(re.escape(k) for k in sorted(replacements, key=len, reverse=True)))

    def resolve(match: re.Match) -> str:
        key = match.group(0)
        value = replacements[key]
        if isinstance(value, list):
            idx = counters[key]
            counters[key] += 1
            if idx >= len(value):
                raise ValueError(f"Недостатньо значень для плейсхолдера {key!r}")
            return value[idx]
        return value

    for paragraph in _iter_all_paragraphs(doc):
        full_text = "".join(r.text for r in paragraph.runs)
        if "<" not in full_text:
            continue
        new_text = pattern.sub(resolve, full_text)
        if new_text != full_text:
            _set_paragraph_text(paragraph, new_text)


def _has_real_numbering(paragraph: Paragraph) -> bool:
    """
    True лише якщо абзац дійсно належить до видимого нумерованого списку.
    Word інколи лишає елемент <w:numPr> на абзацах ПІСЛЯ списку (напр. після
    видалення нумерації), але з <w:numId w:val="0"/> -- це службове значення
    "без нумерації", а не посилання на реальний список, тому такі абзаци
    потрібно відрізняти від справжніх пунктів "НАКАЗУЮ:".
    """
    ppr = paragraph._p.pPr
    if ppr is None:
        return False
    numpr = ppr.find(qn("w:numPr"))
    if numpr is None:
        return False
    numid_el = numpr.find(qn("w:numId"))
    if numid_el is None:
        return False
    numid_val = numid_el.get(qn("w:val"))
    return numid_val is not None and numid_val != "0"


def replace_numbered_clauses(doc: DocumentObject, marker_text: str, clauses: list[str]) -> None:
    """
    Знаходить абзац-маркер (напр. "НАКАЗУЮ:"), бере наступні за ним
    абзаци нумерованого списку (до першої таблиці/кінця нумерованого
    списку) і замінює їх на потрібну кількість пунктів із текстами clauses.
    """
    all_paragraphs = list(doc.paragraphs)

    marker_index = None
    for i, p in enumerate(all_paragraphs):
        if p.text.strip() == marker_text:
            marker_index = i
            break
    if marker_index is None:
        raise ValueError(f"Не знайдено абзац-маркер {marker_text!r} у шаблоні")

    list_paragraphs: list[Paragraph] = []
    for p in all_paragraphs[marker_index + 1 :]:
        if _has_real_numbering(p):
            list_paragraphs.append(p)
            continue
        if p.text.strip() == "":
            # порожній рядок-роздільник перед списком або між пунктами
            continue
        break

    if not list_paragraphs:
        raise ValueError(f"Не знайдено пунктів нумерованого списку після {marker_text!r}")

    template_p = list_paragraphs[0]

    # Додаємо потрібну кількість клонів одразу після останнього наявного пункту.
    anchor = list_paragraphs[-1]._p
    while len(list_paragraphs) < len(clauses):
        new_p_xml = copy.deepcopy(template_p._p)
        anchor.addnext(new_p_xml)
        anchor = new_p_xml
        list_paragraphs.append(Paragraph(new_p_xml, template_p._parent))

    # Зайві пункти (якщо клаузів менше, ніж було в шаблоні) видаляємо.
    while len(list_paragraphs) > len(clauses):
        extra = list_paragraphs.pop()
        extra._p.getparent().remove(extra._p)

    for paragraph, text in zip(list_paragraphs, clauses):
        _set_paragraph_text(paragraph, text)


def remove_table_row_by_first_cell(doc: DocumentObject, first_cell_text: str) -> None:
    """
    Видаляє рядок таблиці за текстом його першої комірки -- використовується,
    щоб прибрати рядок "<директор ДСР або профільний проректор>" з блоку
    ПОГОДЖЕНО (цю роль тепер виконує підписант, вказаний згори наказу).
    """
    for table in doc.tables:
        for row in table.rows:
            if row.cells and row.cells[0].text.strip() == first_cell_text:
                row._tr.getparent().remove(row._tr)
                return


def fill_order_template(
    template_path: Path,
    simple_replacements: dict[str, str | list[str]],
    clauses: list[str],
    naming_marker: str = "НАКАЗУЮ:",
) -> DocumentObject:
    doc = docx.Document(str(template_path))
    replace_numbered_clauses(doc, naming_marker, clauses)
    remove_table_row_by_first_cell(doc, "<директор ДСР або профільний проректор>")
    replace_simple_placeholders(doc, simple_replacements)
    return doc
