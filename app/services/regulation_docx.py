"""Побудова типового Положення про гурток як .docx (Times New Roman)."""

from __future__ import annotations

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from app.db.models import Club, DIRECTION_LABELS
from app.services import dates_uk

FONT_NAME = "Times New Roman"
BASE_SIZE = Pt(14)


def _add_paragraph(doc, text: str = "", bold: bool = False, align=None, size: Pt | None = None) -> None:
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    run.font.name = FONT_NAME
    run.font.size = size or BASE_SIZE
    run.font.bold = bold


def _add_heading(doc, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = FONT_NAME
    run.font.size = BASE_SIZE
    run.font.bold = True


def _add_numbered(doc, items: list[str]) -> None:
    for i, item in enumerate(items, start=1):
        _add_paragraph(doc, f"{i}. {item}")


def build_regulation_document(club: Club) -> docx.document.Document:
    direction_nomn = DIRECTION_LABELS[club.direction]["nomn"]
    direction_plur_gent = DIRECTION_LABELS[club.direction]["plur_gent"]
    faculty_genitive = club.faculty.genitive
    order_reference = dates_uk.format_order_reference(club.order_number, club.order_date)

    doc = docx.Document()
    style = doc.styles["Normal"]
    style.font.name = FONT_NAME
    style.font.size = BASE_SIZE

    _add_paragraph(doc, "ЗАТВЕРДЖЕНО", align=WD_ALIGN_PARAGRAPH.RIGHT)
    _add_paragraph(doc, "наказом КПІ ім. Ігоря Сікорського", align=WD_ALIGN_PARAGRAPH.RIGHT)
    _add_paragraph(doc, order_reference, align=WD_ALIGN_PARAGRAPH.RIGHT)
    _add_paragraph(doc)

    _add_paragraph(
        doc,
        "ПОЛОЖЕННЯ",
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        size=Pt(16),
    )
    _add_paragraph(
        doc,
        f"про гурток «{club.name}» {direction_nomn} спрямування",
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        size=Pt(16),
    )
    _add_paragraph(doc)

    _add_heading(doc, "1. Загальні положення")
    _add_numbered(
        doc,
        [
            (
                f"Гурток «{club.name}» {direction_nomn} спрямування (далі — Гурток) є добровільним "
                f"об’єднанням здобувачів вищої освіти КПІ ім. Ігоря Сікорського, створеним з метою "
                f"розвитку {direction_plur_gent} здібностей його учасників."
            ),
            (
                f"Гурток діє на базі {faculty_genitive} КПІ ім. Ігоря Сікорського за організаційної "
                f"підтримки Департаменту студентського розвитку (відділу розвитку студентського "
                f"потенціалу)."
            ),
            (
                "У своїй діяльності Гурток керується законодавством України, Статутом "
                "КПІ ім. Ігоря Сікорського, наказами і розпорядженнями університету та цим Положенням."
            ),
        ],
    )
    _add_paragraph(doc)

    _add_heading(doc, "2. Мета та завдання")
    _add_numbered(
        doc,
        [
            (
                "Метою діяльності Гуртка є створення умов для розвитку творчого, інтелектуального "
                "та практичного потенціалу здобувачів вищої освіти за обраним спрямуванням."
            ),
            (
                "Основними завданнями Гуртка є: залучення здобувачів вищої освіти до систематичної "
                "позанавчальної діяльності за спрямуванням Гуртка; підвищення рівня практичних "
                "навичок учасників; організація та участь у заходах, конкурсах, змаганнях, "
                "конференціях відповідного спрямування; популяризація напряму діяльності Гуртка "
                "серед студентської спільноти університету."
            ),
        ],
    )
    _add_paragraph(doc)

    _add_heading(doc, "3. Керівництво Гуртком")
    _add_numbered(
        doc,
        [
            (
                "Загальне керівництво діяльністю Гуртка здійснює керівник (керівники) Гуртка, "
                "призначений(-і) відповідним наказом ректора."
            ),
            (
                "Керівник Гуртка організовує поточну діяльність Гуртка, забезпечує ведення "
                "звітності, взаємодіє з відділом розвитку студентського потенціалу ДСР із питань "
                "організаційного та методичного супроводу Гуртка."
            ),
            "Керівник Гуртка здійснює свої повноваження на громадських засадах, без додаткової оплати праці.",
        ],
    )
    _add_paragraph(doc)

    _add_heading(doc, "4. Члени Гуртка, їхні права та обов’язки")
    _add_numbered(
        doc,
        [
            (
                "Членом Гуртка може бути будь-який здобувач вищої освіти КПІ ім. Ігоря Сікорського, "
                "який виявив бажання брати участь у діяльності Гуртка."
            ),
            (
                "Члени Гуртка мають право: брати участь у заходах, що проводить Гурток; "
                "використовувати матеріально-технічну базу, надану Гуртку для здійснення його "
                "діяльності; вносити пропозиції щодо напрямів і форм роботи Гуртка."
            ),
            (
                "Члени Гуртка зобов’язані: дотримуватися цього Положення; брати активну участь у "
                "діяльності Гуртка; дбайливо ставитися до наданого Гуртку майна."
            ),
        ],
    )
    _add_paragraph(doc)

    _add_heading(doc, "5. Організація діяльності")
    _add_numbered(
        doc,
        [
            "Гурток провадить свою діяльність відповідно до плану роботи, що складається керівником Гуртка на навчальний рік.",
            (
                "Керівник Гуртка щорічно звітує про діяльність Гуртка перед відділом розвитку "
                "студентського потенціалу ДСР у порядку, визначеному ДСР."
            ),
            "Засідання (заняття) Гуртка проводяться регулярно, за розкладом, погодженим керівником Гуртка з учасниками.",
        ],
    )
    _add_paragraph(doc)

    _add_heading(doc, "6. Припинення діяльності")
    _add_numbered(
        doc,
        [
            "Діяльність Гуртка припиняється на підставі відповідного наказу КПІ ім. Ігоря Сікорського за поданням ВРСП.",
            (
                "Підставою для припинення діяльності Гуртка може бути звернення керівника Гуртка, "
                "тривала відсутність активної діяльності Гуртка або інші обґрунтовані причини."
            ),
        ],
    )

    return doc
